from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Side, Font
from pathlib import Path
import calendar


def criar_cartao_ponto(nome, mes, ano, pasta_saida):
    pasta = Path(pasta_saida)
    pasta.mkdir(parents=True, exist_ok=True)

    caminho = pasta / f"{nome}_{mes:02d}_{ano}.xlsx"

    if caminho.exists():
        return load_workbook(caminho), caminho

    wb = Workbook()
    ws = wb.active
    ws.title = "Cartão Ponto"

    # Estilos
    center = Alignment(horizontal="center", vertical="center")
    bold = Font(bold=True)
    borda = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Cabeçalho
    ws.merge_cells("A1:G1")
    ws["A1"] = "CARTÃO PONTO"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = center

    ws["A2"] = "Nome:"
    ws["B2"] = nome
    ws["E2"] = "Mês:"
    ws["F2"] = f"{mes:02d}/{ano}"

    # Títulos
    ws.append([
        "Data",
        "Manhã Entrada",
        "Manhã Saída",
        "Tarde Entrada",
        "Tarde Saída",
        "Total de Horas",
        "Visto"
    ])

    for col in "ABCDEFG":
        ws[f"{col}3"].font = bold
        ws[f"{col}3"].alignment = center
        ws[f"{col}3"].border = borda

    # Linhas do mês
    dias_mes = calendar.monthrange(ano, mes)[1]

    for dia in range(1, dias_mes + 1):
        ws.append([f"{dia:02d}"] + [""] * 6)
        linha = ws.max_row
        for col in "ABCDEFG":
            ws[f"{col}{linha}"].alignment = center
            ws[f"{col}{linha}"].border = borda

    wb.save(caminho)
    return wb, caminho


def preencher_dia(ws, registro, total_horas):
    dia = registro.data.day
    linha = 3 + dia  # cabeçalho + títulos

    ws[f"B{linha}"] = registro.entrada_manha
    ws[f"C{linha}"] = registro.saida_manha
    ws[f"D{linha}"] = registro.entrada_tarde
    ws[f"E{linha}"] = registro.saida_tarde
    ws[f"F{linha}"] = total_horas
